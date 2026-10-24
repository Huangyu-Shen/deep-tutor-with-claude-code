"""Build the small, self-authored test fixtures (no third-party material).

Development-time only:
    python3 tests/fixtures/build.py <path-to-python-with-python-pptx>

Creates in tests/fixtures/:
  lecture.pdf  3-page PDF with a text layer (hand-written, stdlib only)
  scanned.pdf  2-page image-only PDF (no text layer; Pillow)
  slides.pptx  3 slides (python-pptx, run in a second interpreter)
  notes.md     plain Markdown
"""

from __future__ import annotations

from pathlib import Path
import subprocess
import sys

HERE = Path(__file__).resolve().parent

LECTURE_PAGES = [
    ["Lecture 1: Photosynthesis", "The light reactions occur in the thylakoid membranes.",
     "Chlorophyll absorbs light mainly in the blue and red wavelengths."],
    ["The Calvin cycle", "The Calvin cycle fixes carbon dioxide into sugar.",
     "Rubisco is the enzyme that catalyses carbon fixation."],
    ["Summary", "Photosynthesis converts light energy into chemical energy.",
     "Oxygen is released as a by-product of water splitting."],
]
SCANNED_PAGES = [
    ["Scanned page one", "Mitochondria produce ATP through oxidative phosphorylation."],
    ["Scanned page two", "The electron transport chain pumps protons across the inner membrane."],
]


def pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_text_pdf(pages: list[list[str]]) -> bytes:
    objs: list[bytes] = []
    n = len(pages)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, lines in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        ops = ["BT", "/F1 14 Tf", "72 720 Td", "20 TL"]
        for line in lines:
            ops.append(f"({pdf_escape(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("cp1252")
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def build_two_column_pdf(left: list[str], right: list[str]) -> bytes:
    """One page with two text columns whose lines sit at the same heights.

    pdftotext -layout puts the two columns side by side line by line, so a sentence that runs over
    two left-column lines comes out interleaved with the right column (as on real two-column slides).
    """
    ops = ["BT", "/F1 10 Tf"]
    for column, x in ((left, 72), (right, 330)):
        for i, line in enumerate(column):
            ops.append(f"1 0 0 1 {x} {700 - 12 * i} Tm ({pdf_escape(line)}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode("cp1252")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def build_scanned_pdf(pages: list[list[str]], path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 34)
    except OSError:
        font = ImageFont.load_default()
    images = []
    for lines in pages:
        img = Image.new("RGB", (1240, 1754), "white")
        draw = ImageDraw.Draw(img)
        for i, line in enumerate(lines):
            draw.text((100, 160 + i * 70), line, fill="black", font=font)
        images.append(img)
    images[0].save(path, "PDF", resolution=150, save_all=True, append_images=images[1:])


PPTX_SCRIPT = r"""
import sys
from pptx import Presentation
prs = Presentation()
slides = [
    ("Cell respiration", ["Glycolysis happens in the cytoplasm.", "It yields two ATP per glucose."]),
    ("Krebs cycle", ["The Krebs cycle takes place in the mitochondrial matrix."]),
    ("Fermentation", ["Without oxygen, cells regenerate NAD+ by fermentation."]),
]
for title, bullets in slides:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = "\n".join(bullets)
prs.save(sys.argv[1])
"""


def main() -> None:
    (HERE / "lecture.pdf").write_bytes(build_text_pdf(LECTURE_PAGES))
    build_scanned_pdf(SCANNED_PAGES, HERE / "scanned.pdf")
    (HERE / "notes.md").write_text(
        "# Study notes\n\nOsmosis is the diffusion of water across a semipermeable membrane.\n",
        encoding="utf-8",
    )
    if len(sys.argv) > 1:
        subprocess.run([sys.argv[1], "-c", PPTX_SCRIPT, str(HERE / "slides.pptx")], check=True)
    print("fixtures written to", HERE)


if __name__ == "__main__":
    main()
